"""双击打开的桌面窗口。"""

from __future__ import annotations

import json
import re
import shutil
import sqlite3
from pathlib import Path

import webview
from webview.dom import DOMEventHandler

import app

OPEN_TYPES = (
    "数据文件 (*.xlsx;*.xlsm;*.xls;*.csv;*.json;*.db;*.sqlite;*.sqlite3)",
    "Excel (*.xlsx;*.xlsm;*.xls)",
    "CSV (*.csv)",
    "JSON (*.json)",
    "SQLite (*.db;*.sqlite;*.sqlite3)",
)
CSV_TYPES = ("CSV (*.csv)",)
XLSX_TYPES = ("Excel (*.xlsx)",)
JSON_TYPES = ("JSON (*.json)",)
DB_TYPES = ("SQLite (*.db)",)


class Api:
    def schema(self):
        return {"tables": app.list_tables()}

    def query(self, sql: str):
        try:
            return app.run_sql(sql)
        except app.AppError as exc:
            return {"error": str(exc)}

    def reset_data(self):
        return app.reset()

    def choose_files(self):
        selected = webview.windows[0].create_file_dialog(
            webview.FileDialog.OPEN,
            allow_multiple=True,
            file_types=OPEN_TYPES,
        )
        if not selected:
            return None
        try:
            return app.import_paths(list(selected))
        except app.AppError as exc:
            return {"error": str(exc)}

    def save_csv(self, sql: str):
        return _save_query(sql, "query.csv", CSV_TYPES, ".csv", write_csv)

    def save_excel(self, sql: str):
        return _save_query(sql, "query.xlsx", XLSX_TYPES, ".xlsx", write_excel)

    def save_json(self, sql: str):
        return _save_query(sql, "query.json", JSON_TYPES, ".json", write_json)

    def save_db(self, sql: str):
        selected = _choose_save("query.db", DB_TYPES)
        if not selected:
            return None
        path = _with_suffix(selected, ".db", app.SQLITE_EXTS)
        if app._same_file(path, app.DB_PATH):
            return {"error": "不能覆盖当前正在使用的数据库"}
        return _write_query(sql, path, write_sqlite)


def _choose_save(filename: str, file_types):
    return webview.windows[0].create_file_dialog(
        webview.FileDialog.SAVE,
        save_filename=filename,
        file_types=file_types,
    )


def _with_suffix(selected, suffix: str, allowed: set[str] | None = None) -> Path:
    path = Path(selected if isinstance(selected, str) else selected[0])
    if allowed is None:
        if path.suffix.lower() != suffix:
            path = path.with_suffix(suffix)
    elif path.suffix.lower() not in allowed:
        path = path.with_suffix(suffix)
    return path


def _save_query(sql: str, filename: str, file_types, suffix: str, writer) -> str | dict | None:
    selected = _choose_save(filename, file_types)
    if not selected:
        return None
    return _write_query(sql, _with_suffix(selected, suffix), writer)


def _write_query(sql: str, path: Path, writer) -> str | dict:
    try:
        with app.query_rows(sql) as (columns, rows):
            writer(path, columns, rows)
    except app.AppError as exc:
        return {"error": str(exc)}
    except Exception as exc:
        return {"error": f"无法保存：{exc}"}
    return str(path)


def _fit_row(row, width: int) -> list:
    values = list(row)
    if len(values) < width:
        values.extend([None] * (width - len(values)))
    return values[:width]


def csv_cell(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if any(char in text for char in '",\n\r'):
        return '"' + text.replace('"', '""') + '"'
    return text


def write_csv(path: Path, columns, rows) -> None:
    names = app.dedupe_names([str(column) for column in columns])
    lines = [",".join(csv_cell(name) for name in names)]
    width = len(names)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        handle.write(lines[0])
        for row in rows:
            handle.write("\r\n")
            handle.write(",".join(csv_cell(value) for value in _fit_row(row, width)))
        handle.write("\r\n")


def write_json(path: Path, columns, rows) -> None:
    names = app.dedupe_names([str(column) for column in columns])
    width = len(names)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write("[\n")
        first = True
        for row in rows:
            if not first:
                handle.write(",\n")
            first = False
            values = _fit_row(row, width)
            json.dump(
                {name: values[index] for index, name in enumerate(names)},
                handle,
                ensure_ascii=False,
            )
        handle.write("\n]\n")


def write_sqlite(path: Path, columns, rows) -> None:
    for extra in (path, Path(str(path) + "-wal"), Path(str(path) + "-shm")):
        if extra.is_file():
            extra.unlink()
    names = app.dedupe_names([str(column) for column in columns])
    column_sql = ", ".join(app.quote_ident(name) for name in names)
    conn = sqlite3.connect(path)
    try:
        conn.execute(f"CREATE TABLE {app.quote_ident('查询结果')} ({column_sql})")
        placeholders = ", ".join(["?"] * len(names))
        batch: list[list] = []
        sql = f"INSERT INTO {app.quote_ident('查询结果')} ({column_sql}) VALUES ({placeholders})"
        for row in rows:
            batch.append(_fit_row(row, len(names)))
            if len(batch) >= 1000:
                conn.executemany(sql, batch)
                batch.clear()
        if batch:
            conn.executemany(sql, batch)
        conn.commit()
    finally:
        conn.close()


def write_excel(path: Path, columns, rows) -> None:
    from openpyxl import Workbook

    book = Workbook(write_only=True)
    sheet = book.create_sheet("查询结果")
    names = [str(column) for column in columns]
    sheet.append([_excel_cell(column) for column in names])
    width = len(names)
    for row in rows:
        sheet.append([_excel_cell(value) for value in _fit_row(row, width)])
    book.save(path)


def _excel_cell(value):
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    return re.sub(r"[\000-\010\013\014\016-\037]", "", text)


def dropped_paths(event: dict) -> list[str]:
    files = (event.get("dataTransfer") or {}).get("files") or []
    paths: list[str] = []
    for item in files:
        if not isinstance(item, dict):
            continue
        path = item.get("pywebviewFullPath") or ""
        if path:
            paths.append(path)
    return paths


def on_drop(window, event: dict):
    paths = dropped_paths(event)
    if not paths:
        window.evaluate_js("window.onImportFinished && window.onImportFinished()")
        return
    window.evaluate_js("window.onImportStart && window.onImportStart()")
    try:
        payload = app.import_paths(paths)
    except app.AppError as exc:
        payload = {"error": str(exc)}
    window.evaluate_js(f"window.onImported({json.dumps(payload, ensure_ascii=False)})")


def bind(window):
    window.dom.document.events.drop += DOMEventHandler(lambda event: on_drop(window, event), True, True)


def clear_page_cache() -> None:
    """WebView2 会把固定地址的页面留在缓存里，改过的页面因此打不开。只删缓存，保留本地记录。"""
    root = app.DATA_DIR / "webview"
    if not root.exists():
        return
    cache_names = {
        "Cache",
        "Code Cache",
        "GPUCache",
        "DawnGraphiteCache",
        "DawnWebGPUCache",
        "ShaderCache",
        "GrShaderCache",
    }
    targets = [path for path in root.rglob("*") if path.is_dir() and path.name in cache_names]
    for path in targets:
        shutil.rmtree(path, ignore_errors=True)


def main():
    window = webview.create_window(
        "SQL 查 Excel",
        url=str(app.STATIC_DIR / "index.html"),
        js_api=Api(),
        width=1200,
        height=800,
        min_size=(960, 640),
        text_select=True,
        background_color="#E7E1D6",
    )
    app.DATA_DIR.mkdir(parents=True, exist_ok=True)
    clear_page_cache()
    webview.start(bind, (window,), private_mode=False, storage_path=str(app.DATA_DIR / "webview"))


if __name__ == "__main__":
    main()
